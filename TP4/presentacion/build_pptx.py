#!/usr/bin/env python3
"""Arma el .pptx que se usa PARA PRESENTAR, con las animaciones embebidas.

El PDF que se entrega (modo `entrega`) lleva un fotograma fijo y el link de cada
animacion (YouTube o Vimeo); la guia (2.4.8) pide que durante la exposicion en vivo
las animaciones se reproduzcan en la diapositiva. PowerPoint reproduce GIF animados
en modo presentacion, asi que la version en vivo es un .pptx.

Usar el PDF de `entrega` para el campus y este .pptx SOLO para la exposicion: el
.pptx no se entrega nunca y esta en .gitignore, igual que los MP4 y los GIF.

El .pptx no se re-maqueta a mano: se compila `presentacion.tex` en modo `vivo`
(cada `\\animacion{<id>}{...}` deja un cuadrado magenta en vez del fotograma), se
rasteriza cada pagina y se pega esa imagen a sangre en una diapositiva. Encima de
cada cuadrado, localizado por color y no por coordenadas escritas a mano, va el GIF.
El orden de los GIF sale del propio `presentacion.tex` (`animation_sequence`): cada
llamada a `\\animacion` consume un hueco, de izquierda a derecha y de arriba a abajo.

Los GIF se generan desde los MP4 de `presentacion/animaciones/` con ffmpeg (paleta de
64 colores por video, 640 px, 15 fps) y se cachean en `animaciones/gif_pptx/`.
PyMuPDF y python-pptx (solo en el Python de Windows) se importan recien al armar el
.pptx, de modo que el resto del modulo y sus tests no los necesitan.

    py -3.14 presentacion/build_pptx.py            # desde TP4/
    py -3.14 presentacion/build_pptx.py --dpi 220  # mas resolucion, .pptx mas pesado
"""

from __future__ import annotations

import argparse
import re
import subprocess
import sys
from pathlib import Path

import numpy as np
from PIL import Image

PRES_DIR = Path(__file__).resolve().parent
TP4_DIR = PRES_DIR.parent
sys.path.insert(0, str(PRES_DIR))

import links  # noqa: E402

TEX_NAME = "presentacion.tex"
LIVE_JOBNAME = "presentacion_vivo"
VIDEO_DIR = PRES_DIR / "animaciones"
GIF_DIR = VIDEO_DIR / "gif_pptx"
GIF_WIDTH = 640
GIF_FPS = 15

# Mismo valor que el cuadrado de `\animacion` en modo vivo del tex.
MARKER_RGB = (255, 0, 255)
MARKER_TOL = 12

# Diapositiva 16:9 (la misma relacion que `aspectratio=169`), en EMU.
EMU_PER_INCH = 914400
SLIDE_W = int(round(13.333 * EMU_PER_INCH))
SLIDE_H = int(round(7.5 * EMU_PER_INCH))

_DEFINITION = re.compile(
    r"\\(?:re)?newcommand\*?\s*\{?\\animacion\}?|\\providecommand\*?\s*\{?\\animacion\}?"
    r"|\\DeclareRobustCommand\*?\s*\{?\\animacion\}?|\\def\\animacion(?![A-Za-z])"
)
_CALL = re.compile(r"\\animacion\s*\{([^{}]*)\}")


def _strip_comments(text: str) -> str:
    """Quita lo que sigue a un % sin escapar en cada linea."""
    out = []
    for line in text.splitlines():
        i = 0
        while i < len(line):
            if line[i] == "\\":
                i += 2
                continue
            if line[i] == "%":
                line = line[:i]
                break
            i += 1
        out.append(line)
    return "\n".join(out)


def _skip_definitions(text: str) -> str:
    """Quita la definicion del macro \\animacion (cuerpo de llaves balanceadas)."""
    while True:
        match = _DEFINITION.search(text)
        if match is None:
            return text
        pos = match.end()
        # Argumentos opcionales ([n], #1#2) y despues un grupo de llaves balanceado.
        while pos < len(text) and text[pos] != "{":
            pos += 1
        depth = 0
        while pos < len(text):
            if text[pos] == "\\":
                pos += 2
                continue
            if text[pos] == "{":
                depth += 1
            elif text[pos] == "}":
                depth -= 1
                if depth == 0:
                    pos += 1
                    break
            pos += 1
        text = text[:match.start()] + text[pos:]


def animation_sequence(tex_text: str) -> list[str]:
    """Ids de animacion en el orden en que el tex llama a \\animacion."""
    text = _skip_definitions(_strip_comments(tex_text))
    ids = []
    for found in _CALL.findall(text):
        case_id = found.strip()
        if case_id not in links.CASE_IDS:
            raise ValueError(
                f"\\animacion{{{case_id}}}: id desconocido (validos: {', '.join(links.CASE_IDS)})")
        ids.append(case_id)
    return ids


def ensure_gifs(ids, force: bool = False) -> list[Path]:
    """Convierte a GIF (solo para el .pptx en vivo) cada MP4 que lo necesite."""
    GIF_DIR.mkdir(parents=True, exist_ok=True)
    out = []
    for case_id in dict.fromkeys(ids):
        mp4 = VIDEO_DIR / f"{case_id}.mp4"
        gif = GIF_DIR / f"{case_id}.gif"
        if not mp4.is_file():
            raise FileNotFoundError(
                f"falta {mp4}: correr presentacion/make_animations.py para generar los videos")
        out.append(gif)
        if not force and gif.is_file() and gif.stat().st_mtime >= mp4.stat().st_mtime:
            continue
        vf = (f"fps={GIF_FPS},scale={GIF_WIDTH}:-1:flags=lanczos,split[a][b];"
              "[a]palettegen=max_colors=64:stats_mode=diff[p];"
              "[b][p]paletteuse=dither=none:diff_mode=rectangle")
        subprocess.run(["ffmpeg", "-y", "-v", "error", "-i", str(mp4), "-vf", vf,
                        "-loop", "0", str(gif)], check=True)
        print(f"{gif.name}: {gif.stat().st_size / 1e6:.1f} MB")
    return out


def compile_live_pdf(passes: int = 2) -> Path:
    """Compila la variante `vivo`; dos pasadas porque miniframes necesita el .aux previo."""
    for i in range(passes):
        proc = subprocess.run(
            ["pdflatex", "-interaction=nonstopmode", f"-jobname={LIVE_JOBNAME}",
             rf"\def\modo{{vivo}}\input{{{TEX_NAME}}}"],
            cwd=PRES_DIR, capture_output=True, text=True, check=False,
        )
        if proc.returncode != 0:
            tail = "\n".join(proc.stdout.splitlines()[-25:])
            raise RuntimeError(f"pdflatex fallo en la pasada {i + 1}:\n{tail}")
    return PRES_DIR / f"{LIVE_JOBNAME}.pdf"


def marker_mask(page_rgb: np.ndarray) -> np.ndarray:
    """Mascara booleana de los pixeles del color marcador."""
    return np.all(np.abs(page_rgb.astype(int) - MARKER_RGB) <= MARKER_TOL, axis=2)


def marker_boxes(mask: np.ndarray) -> list[tuple[int, int, int, int]]:
    """Cuadrados marcadores (x0, y0, x1, y1) en pixeles, de izquierda a derecha.

    Se separan por columnas de pixeles marcados: los huecos de una diapositiva estan
    siempre uno al lado del otro, nunca apilados.
    """
    if not mask.any():
        return []
    columns = np.nonzero(mask.any(axis=0))[0]
    runs = []
    start = prev = int(columns[0])
    for col in columns[1:]:
        col = int(col)
        if col > prev + 1:
            runs.append((start, prev))
            start = col
        prev = col
    runs.append((start, prev))
    out = []
    for x0, x1 in runs:
        rows = np.nonzero(mask[:, x0:x1 + 1].any(axis=1))[0]
        out.append((x0, int(rows[0]), x1, int(rows[-1])))
    return out


def fit_box(box_px, page_px, image_path: Path):
    """(left, top, width, height) en EMU del GIF centrado en el hueco, sin estirarlo."""
    x0, y0, x1, y1 = box_px
    page_w, page_h = page_px
    box_w = (x1 - x0 + 1) / page_w * SLIDE_W
    box_h = (y1 - y0 + 1) / page_h * SLIDE_H
    left = x0 / page_w * SLIDE_W
    top = y0 / page_h * SLIDE_H
    with Image.open(image_path) as img:
        aspect = img.width / img.height
    if box_w / box_h > aspect:
        height, width = box_h, box_h * aspect
    else:
        width, height = box_w, box_w / aspect
    left += (box_w - width) / 2
    top += (box_h - height) / 2
    return int(left), int(top), int(width), int(height)


def build(pdf_path: Path, out_path: Path, dpi: int, page_dir: Path, sequence) -> int:
    """Una diapositiva por pagina del PDF vivo, con cada GIF sobre su hueco."""
    import fitz  # PyMuPDF (solo Windows)
    from pptx import Presentation
    from pptx.util import Emu

    sequence = list(sequence)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(SLIDE_W), Emu(SLIDE_H)
    blank = prs.slide_layouts[6]
    page_dir.mkdir(parents=True, exist_ok=True)
    doc = fitz.open(pdf_path)
    pending = list(sequence)
    holes = 0
    placed = 0

    for index, page in enumerate(doc):
        pix = page.get_pixmap(dpi=dpi)
        rgb = (np.frombuffer(pix.samples, dtype=np.uint8)
               .reshape(pix.height, pix.width, pix.n)[:, :, :3]).copy()
        boxes = marker_boxes(marker_mask(rgb))
        # Se blanquea el cuadrado entero con margen: el borde antialiaseado no pasa el
        # filtro de color y dejaria un contorno alrededor de la animacion.
        margin = max(2, round(dpi / 50))
        for x0, y0, x1, y1 in boxes:
            rgb[max(0, y0 - margin):y1 + 1 + margin, max(0, x0 - margin):x1 + 1 + margin] = 255
        png = page_dir / f"pagina{index + 1:02d}.png"
        Image.fromarray(rgb).save(png)

        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(png), 0, 0, width=Emu(SLIDE_W), height=Emu(SLIDE_H))
        for box in boxes:
            holes += 1
            if not pending:
                continue
            gif = GIF_DIR / f"{pending.pop(0)}.gif"
            if not gif.is_file():
                raise FileNotFoundError(f"falta {gif}: correr ensure_gifs primero")
            left, top, width, height = fit_box(box, (pix.width, pix.height), gif)
            slide.shapes.add_picture(str(gif), Emu(left), Emu(top), width=Emu(width),
                                     height=Emu(height))
            placed += 1

    if holes != len(sequence):
        raise RuntimeError(
            f"el PDF vivo tiene {holes} huecos de animacion y presentacion.tex llama "
            f"{len(sequence)} veces a \\animacion")
    try:
        prs.save(out_path)
    except PermissionError:
        # Windows bloquea el archivo mientras PowerPoint lo tiene abierto.
        raise SystemExit(
            f"error: {out_path} esta abierto en PowerPoint (o bloqueado por otro "
            f"proceso). Cerralo y volve a correr."
        )
    return placed


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--dpi", type=int, default=200,
                        help="resolucion del rasterizado de cada pagina (default 200)")
    parser.add_argument("--out", type=Path, default=PRES_DIR / "presentacion_vivo.pptx")
    parser.add_argument("--skip-latex", action="store_true",
                        help="reusar presentacion_vivo.pdf en vez de recompilarlo")
    args = parser.parse_args(argv)

    tex = PRES_DIR / TEX_NAME
    if not tex.is_file():
        print(f"error: no existe {tex}")
        return 1
    try:
        sequence = animation_sequence(tex.read_text(encoding="utf-8"))
        ensure_gifs(sequence)
        pdf = PRES_DIR / f"{LIVE_JOBNAME}.pdf"
        if not args.skip_latex:
            pdf = compile_live_pdf()
        elif not pdf.is_file():
            print(f"error: no existe {pdf} y se pidio --skip-latex")
            return 1
        placed = build(pdf, args.out, args.dpi, PRES_DIR / "_paginas_vivo", sequence)
    except (ValueError, FileNotFoundError, RuntimeError, subprocess.CalledProcessError,
            ImportError) as exc:
        print(f"error: {exc}")
        return 1
    print(f"{args.out}  ({placed} animaciones embebidas, {args.out.stat().st_size / 1e6:.1f} MB)")
    print("Abrirlo con PowerPoint y presentar con F5: los GIF se reproducen solos.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
